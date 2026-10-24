"""Printable PDF + editable Excel for one old part, grouped block -> house number.
   python3 okhla_import/part_report.py 65   (reads the local D1 copy; writes okhla_import/out/)"""
import sqlite3, glob, sys, datetime, csv, re
from collections import Counter
from collections import OrderedDict
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Table, TableStyle, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PART = int(sys.argv[1]) if len(sys.argv) > 1 else 65
OUT = "okhla_import/out"
F = "/System/Library/Fonts/Supplemental/"
pdfmetrics.registerFont(TTFont("A", F + "Arial.ttf")); pdfmetrics.registerFont(TTFont("AB", F + "Arial Bold.ttf"))

db = [x for x in glob.glob(".wrangler/state/v3/d1/**/*.sqlite", recursive=True) if not x.endswith("metadata.sqlite")][0]
c = sqlite3.connect(db)
rows = c.execute("""SELECT section_no, section_name, house_raw, hh_key, hh_sort, name, relation_type, relation_name, age, gender,
                           epic_id, status, status_detail, conflict, new_part, serial_2025, house_floor
                    FROM ok_electors WHERE old_part=? ORDER BY (section_no=''), CAST(section_no AS INT), hh_sort, serial_2025""", (PART,)).fetchall()
# how each house number was settled (okhla_import/house_fix.py, written by build_okhla.py)
METHOD = {}
for r in csv.DictReader(open(f"{OUT}/house_methods.csv", encoding="utf-8")):
    if r["old_part"] == str(PART): METHOD[r["serial_2025"]] = r["method"]
CHECK = {"neighbours": "Inferred", "unverified": "Check"}
def house_of(r): return (r[2] or "").split(",")[0].strip()
def floor_of(r): return (r[16] or "").strip()
FLAG = {"ASDD": "ASDD", "NOTICE": "NOTICE", "DELETED": "PRE-SIR DELETION", "NO_SIR": "NOT IN SIR"}
ORDER = {"ASDD": 0, "NOTICE": 1, "DELETED": 2, "NO_SIR": 3}
COL = {"ASDD": "FFE0B2", "NOTICE": "FFF59D", "DELETED": "F8BBD0", "NO_SIR": "CFD8DC"}
houses = OrderedDict(); cleared = {}
for r in rows:
    k = (r[0], r[3])
    if r[11] in FLAG: houses.setdefault(k, []).append(r)
    elif r[11] == "CLEARED": cleared[k] = cleared.get(k, 0) + 1
for k in houses: houses[k].sort(key=lambda r: ORDER[r[11]])
cnt = {s: sum(1 for h in houses.values() for r in h if r[11] == s) for s in FLAG}
tot = sum(cnt.values()); today = datetime.date.today().strftime("%d %b %Y")
def blk(n):
    if not n: return "Added after Jan 2025 — no house number on record"
    return " ".join(n.replace("BLOCK-", "Block").replace(",", " ").split())
def newp(r): return str(r[14]) if r[14] else ""
def detail(r):
    d = r[12] or ""
    return {"ASDD": f"ASDD: {d}", "NOTICE": f"Notice: {d}", "DELETED": d, "NO_SIR": "Not in Draft SIR, ASDD, Notice or deletion lists"}[r[11]]

# ---------- PDF ----------
HEAD = ["House No.", "Floor", "Elector", "Relative", "Age / Sex", "EPIC", "Status", "Detail / reason", "New\npart", "House\ncheck", "Field\nnote"]
W = [22, 17, 42, 36, 14, 23, 23, 56, 10, 15, 19]            # mm, landscape A4 minus margins
W = [w * mm for w in W]
sm = ParagraphStyle("s", fontName="A", fontSize=7.2, leading=8.6)
sb = ParagraphStyle("sb", parent=sm, fontName="AB")
hd = ParagraphStyle("h", parent=sm, fontName="AB", textColor=colors.white, fontSize=7.4)
def P(t, st=sm): return Paragraph(str(t).replace("&", "&amp;").replace("<", "&lt;").replace("\n", "<br/>"), st)
data = [[P(h, hd) for h in HEAD]]; styles = []
first_row = {}
cur_blk = None
for (sec, hk), mem in houses.items():
    r0 = mem[0]
    if sec != cur_blk:
        cur_blk = sec
        data.append([P(blk(r0[1]), ParagraphStyle("b", parent=sb, fontSize=8.5))] + [""] * 10)
        i = len(data) - 1
        styles += [("SPAN", (0, i), (-1, i)), ("BACKGROUND", (0, i), (-1, i), colors.HexColor("#E3E8EF"))]
    start = len(data)
    for j, r in enumerate(mem):
        hcell = ""
        if j == 0:
            n = cleared.get((sec, hk), 0)
            hcell = [P(Counter(house_of(x) for x in mem).most_common(1)[0][0] or "—", sb)] + ([P(f"+{n} cleared", ParagraphStyle("c", parent=sm, textColor=colors.grey, fontSize=6.3))] if n else [])
        data.append([hcell, P(floor_of(r)), P(r[5], sb), P(f"{r[6]}: {r[7]}" if r[7] else ""), P(" / ".join(x for x in (r[8] or "", (r[9] or "")[:1]) if x)), P(r[10] or ""),
                     P(FLAG[r[11]], sb), P(detail(r)), P(newp(r)), P(CHECK.get(METHOD.get(str(r[15])), "")), ""])
        i = len(data) - 1
        styles.append(("BACKGROUND", (6, i), (6, i), colors.HexColor("#" + COL[r[11]])))
        if CHECK.get(METHOD.get(str(r[15]))): styles.append(("BACKGROUND", (9, i), (9, i), colors.HexColor("#FFE0E0")))
    styles += [("LINEABOVE", (0, start), (-1, start), 0.8, colors.HexColor("#555555"))]
    styles.append(("NOSPLIT", (0, start), (-1, len(data) - 1)))

t = Table(data, colWidths=W, repeatRows=1)
t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243B55")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                       ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
                       ("LEFTPADDING", (0, 0), (-1, -1), 2.5), ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
                       ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)] + styles))
def footer(cv, doc):
    cv.saveState(); cv.setFont("A", 7); cv.setFillColor(colors.grey)
    cv.drawString(10 * mm, 7 * mm, f"Okhla (AC 54) · Old Part {PART} · House-wise audit list · Confidential: contains voter personal data · {today}")
    cv.drawRightString(landscape(A4)[0] - 10 * mm, 7 * mm, f"Page {doc.page}"); cv.restoreState()
fn = f"{OUT}/Okhla_Part{PART}_house_wise_audit_list"
doc = BaseDocTemplate(fn + ".pdf", pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm,
                      title=f"Okhla Part {PART} house-wise audit list")
doc.addPageTemplates([PageTemplate(frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, 0, 0, 0, 0)], onPage=footer)])
title = ParagraphStyle("t", fontName="AB", fontSize=13, leading=16)
summ = (f"{tot} electors in {len(houses)} houses — ASDD {cnt['ASDD']} · Notice {cnt['NOTICE']} · Pre-SIR deletion {cnt['DELETED']} · Not in SIR {cnt['NO_SIR']}.  "
        f"Electors cleared in the Draft SIR are not listed; “+n cleared” under a house number shows how many more are in that house.  "
        f"House numbers were re-read against the printed roll and two OCR sources; “Inferred” = unreadable card (DELETED stamp), house taken from the cards before/after it in the roll; “Check” = could not be confirmed, please verify at the door.")
doc.build([Paragraph(f"Okhla AC 54 — Old Part {PART}: house-wise list of ASDD, Notice, Pre-SIR deletion and Not-in-SIR electors", title),
           Spacer(1, 2 * mm), Paragraph(summ, ParagraphStyle("u", parent=sm, fontSize=8, leading=10)), Spacer(1, 3 * mm), t])

# ---------- Excel ----------
wb = Workbook(); ws = wb.active; ws.title = f"Part {PART}"
H = ["Block", "House No.", "Floor", "House check", "Elector", "Relation", "Relative name", "Age", "Sex", "EPIC", "Status", "Detail / reason", "Earlier record (note)",
     "Cleared in house (not listed)", "Old part", "New part", "Roll serial 2025", "Field note"]
ws.append(H)
thin = Side(style="thin", color="BBBBBB")
for (sec, hk), mem in houses.items():
    for r in mem:
        ws.append([blk(r[1]), house_of(r), floor_of(r), {"neighbours": "Inferred from neighbours", "unverified": "Check on ground"}.get(METHOD.get(str(r[15])), ""), r[5], r[6], r[7], int(r[8]) if str(r[8] or "").isdigit() else r[8], r[9], r[10], FLAG[r[11]], detail(r),
                   r[13] or "", cleared.get((sec, hk), 0), PART, r[14], r[15], ""])
        ws.cell(ws.max_row, 11).fill = PatternFill("solid", fgColor=COL[r[11]])
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="243B55"); cell.alignment = Alignment(wrap_text=True, vertical="center")
for i, w in enumerate([16, 14, 12, 16, 24, 10, 24, 6, 8, 14, 18, 44, 40, 12, 8, 8, 10, 30], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "D2"; ws.auto_filter.ref = ws.dimensions
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (12, 13, 18)); cell.border = Border(bottom=thin)
ws.print_title_rows = "1:1"; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(fn + ".xlsx")
print(fn, tot, len(houses), cnt)
