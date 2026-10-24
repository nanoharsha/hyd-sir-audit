"""Household-wise printable list (PDF + Excel) of one block: SIR Cleared, ASDD and Notice electors, house by house.
   python3 okhla_import/household_list.py 65 K"""
import sqlite3, glob, sys, datetime, json, subprocess
from collections import OrderedDict, Counter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.graphics.shapes import Drawing, Circle, Line
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Table, TableStyle, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PART, BLOCK = int(sys.argv[1]), sys.argv[2].upper()
KEEP = ("CLEARED", "ASDD", "NOTICE")
FLAG = {"CLEARED": "SIR CLEARED", "ASDD": "ASDD", "NOTICE": "NOTICE"}
COL = {"CLEARED": "C8E6C9", "ASDD": "FFE0B2", "NOTICE": "FFF59D"}
ORDER = {"NOTICE": 0, "ASDD": 1, "CLEARED": 2}
F = "/System/Library/Fonts/Supplemental/"
pdfmetrics.registerFont(TTFont("AU", F + "Arial Unicode.ttf")); pdfmetrics.registerFont(TTFont("A", F + "Arial.ttf")); pdfmetrics.registerFont(TTFont("AB", F + "Arial Bold.ttf"))
def prod(sql):
    o = subprocess.run(["npx", "wrangler", "d1", "execute", "electoral-app-db", "--remote", "--command", sql, "--json"],
                       stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True).stdout
    j = json.loads(o)
    if not isinstance(j, list): raise SystemExit(f"production query failed: {j}")
    return j[0]["results"]
VIS = {r["voter_key"]: r for r in prod(f"SELECT voter_key, form_type, asdd_correct FROM ok_verification WHERE old_part={PART}")}
VISITS = {k: v["form_type"] for k, v in VIS.items()}   # live: who the field team has covered
FORM = {"ASDD": "asdd", "NOTICE": "notice"}
def covered(r): return FORM.get(r[9]) is not None and VISITS.get(norm(r[8])) == FORM[r[9]]
def disputed(r): return r[9] == "ASDD" and covered(r) and VIS[norm(r[8])].get("asdd_correct") == "no"   # the field team says the ASDD marking is wrong
norm = lambda e: "".join(ch for ch in (e or "").upper() if ch.isalnum())
db = [x for x in glob.glob(".wrangler/state/v3/d1/**/*.sqlite", recursive=True) if not x.endswith("metadata.sqlite")][0]
c = sqlite3.connect(db)
rows = c.execute("""SELECT house_raw, hh_key, hh_sort, name, relation_type, relation_name, age, gender, epic_id, status, status_detail,
                           new_part, house_floor, serial_2025 FROM ok_electors
                    WHERE old_part=? AND block=? AND status IN ('CLEARED','ASDD','NOTICE') ORDER BY hh_sort, serial_2025""", (PART, BLOCK)).fetchall()
houses = OrderedDict()
for r in rows: houses.setdefault(r[1], []).append(r)
for k in houses: houses[k].sort(key=lambda r: (r[12] or "", ORDER[r[9]], r[13] or 0))
cnt = Counter(r[9] for r in rows)
nv = sum(1 for r in rows if covered(r)); nd = sum(1 for r in rows if disputed(r)); nform = sum(1 for r in rows if r[9] in FORM)
other = Counter(x[0] for x in c.execute("SELECT status FROM ok_electors WHERE old_part=? AND block=?", (PART, BLOCK)))
def hn(rs): return (rs[0][0] or "").split(",")[0].strip() or "no house number"
def rel(r): return f"{r[4]}: {r[5]}" if r[5] else ""
def det(r): return {"CLEARED": "In Draft SIR 2026", "ASDD": f"ASDD: {r[10] or ''}", "NOTICE": f"Notice: {r[10] or ''}"}[r[9]]

# ---------------- PDF
sm = ParagraphStyle("s", fontName="A", fontSize=7.8, leading=9.4); sb = ParagraphStyle("sb", parent=sm, fontName="AB")
hd = ParagraphStyle("h", parent=sm, fontName="AB", textColor=colors.white)
def P(t, st=sm): return Paragraph(str(t).replace("&", "&amp;").replace("<", "&lt;"), st)
HEAD = ["Sr.", "House No.", "Floor", "Elector", "Relative", "Age / Sex", "EPIC", "Status", "Detail", "New part", "Verified"]
W = [mm * w for w in (9, 23, 16, 46, 42, 15, 25, 25, 52, 12, 14)]
def cross_circle(size=11):
    d = Drawing(size + 2, size + 2); c = size / 2 + 1; r = size / 2
    red = colors.HexColor("#B71C1C")
    d.add(Circle(c, c, r, strokeColor=red, strokeWidth=1.3, fillColor=None))
    k = r * 0.55
    d.add(Line(c - k, c - k, c + k, c + k, strokeColor=red, strokeWidth=1.4)); d.add(Line(c - k, c + k, c + k, c - k, strokeColor=red, strokeWidth=1.4))
    d.hAlign = "CENTER"
    return d
tick = ParagraphStyle("t", fontName="AU", fontSize=10, leading=9.4, alignment=1, textColor=colors.HexColor("#1B7A3E"))
data = [[P(h, hd) for h in HEAD]]; styles = []; sn = 0
for k, rs in houses.items():
    start = len(data); styles.append(("LINEABOVE", (0, start), (-1, start), 0.8, colors.HexColor("#555555")))
    for j, r in enumerate(rs):
        sn += 1
        data.append([P(sn), P(hn(rs) if j == 0 else "", sb), P((r[12] or "").strip()), P(r[3], sb), P(rel(r)),
                     P(" / ".join(x for x in (r[6] or "", (r[7] or "")[:1]) if x)), P(r[8] or ""), P(FLAG[r[9]], sb), P(det(r)), P(r[11] or ""), cross_circle() if disputed(r) else Paragraph("✓", tick) if covered(r) else ""])
        if covered(r): styles.append(("BACKGROUND", (10, len(data) - 1), (10, len(data) - 1), colors.HexColor("#FDECEA" if disputed(r) else "#E6F4EA")))
        styles.append(("BACKGROUND", (7, len(data) - 1), (7, len(data) - 1), colors.HexColor("#" + COL[r[9]])))
t = Table(data, colWidths=W, repeatRows=1)
t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243B55")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (10, 0), (10, -1), "CENTER"),
                       ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")), ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
                       ("RIGHTPADDING", (0, 0), (-1, -1), 2.5), ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)] + styles))
stamp = datetime.datetime.now().strftime("%d %b %Y")
def footer(cv, doc):
    cv.saveState(); cv.setFont("A", 7); cv.setFillColor(colors.grey)
    cv.drawString(10 * mm, 7 * mm, f"Okhla (AC 54) · Old Part {PART} · Block {BLOCK} · household-wise list · {stamp} · Confidential: contains voter personal data")
    cv.drawRightString(landscape(A4)[0] - 10 * mm, 7 * mm, f"Page {doc.page}"); cv.restoreState()
base = f"okhla_import/out/Okhla_Part{PART}_Block{BLOCK}_household_list"
doc = BaseDocTemplate(base + ".pdf", pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm, title=f"Okhla Part {PART} Block {BLOCK} household list")
doc.addPageTemplates([PageTemplate(frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, 0, 0, 0, 0)], onPage=footer)])
summ = (f"{len(rows)} electors in {len(houses)} households, house number small to large: <b>{cnt['CLEARED']} SIR Cleared</b>, <b>{cnt['ASDD']} ASDD</b>, <b>{cnt['NOTICE']} Notice</b>. "
        f"<b>Field verification: {nv} of {nform} ASDD and Notice electors covered</b> — a tick means covered; a red circle with a cross means the field team found the ASDD marking WRONG ({nd} electors). SIR Cleared electors have no verification form, so their box stays blank. "
        f"Not listed here: {other.get('DELETED', 0)} pre-SIR deleted and {other.get('NO_SIR', 0)} not-in-SIR electors of this block.")
doc.build([Paragraph(f"Okhla AC 54 — Old Part {PART}, Block {BLOCK}: household-wise list (SIR Cleared · ASDD · Notice)", ParagraphStyle("t", fontName="AB", fontSize=13, leading=16)),
           Spacer(1, 2 * mm), Paragraph(summ, ParagraphStyle("u", parent=sm, fontSize=8.5, leading=11)), Spacer(1, 3 * mm), t])

# ---------------- Excel
wb = openpyxl.Workbook(); ws = wb.active; ws.title = f"Block {BLOCK}"
ws.append(["Sr.", "House No.", "Floor", "Elector", "Relative", "Age", "Sex", "EPIC", "Status", "Detail", "New part", "Verified", "Household (hh_key)"])
for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="243B55")
sn = 0
for k, rs in houses.items():
    for r in rs:
        sn += 1
        ws.append([sn, hn(rs), (r[12] or "").strip(), r[3], rel(r), r[6], (r[7] or "")[:1], r[8], FLAG[r[9]], det(r), r[11], "⊗ ASDD wrong" if disputed(r) else "✓" if covered(r) else "", k])
        ws.cell(row=ws.max_row, column=9).fill = PatternFill("solid", fgColor=COL[r[9]])
for col, w in zip("ABCDEFGHIJKLM", (6, 14, 10, 28, 30, 6, 5, 14, 14, 36, 9, 10, 22)): ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
wb.save(base + ".xlsx")
print(base, len(rows), "electors", len(houses), "houses", dict(cnt), "| not listed:", {k: v for k, v in other.items() if k not in KEEP})
