"""Printable PDF (+ Excel): individual electors whose ASDD marking, or pre-SIR deletion, the field team found INCORRECT.
Block by block, house number small to large, no phone numbers.   python3 okhla_import/discrepancy_electors.py
Electors from the local D1 copy; visits LIVE from production."""
import sqlite3, glob, json, subprocess, datetime
from collections import defaultdict
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import openpyxl
from openpyxl.styles import Font, PatternFill

F = "/System/Library/Fonts/Supplemental/"
pdfmetrics.registerFont(TTFont("A", F + "Arial.ttf")); pdfmetrics.registerFont(TTFont("AB", F + "Arial Bold.ttf"))
def prod(sql):
    o = subprocess.run(["npx", "wrangler", "d1", "execute", "electoral-app-db", "--remote", "--command", sql, "--json"],
                       stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True).stdout
    j = json.loads(o)
    if not isinstance(j, list): raise SystemExit(f"production query failed: {j}")
    return j[0]["results"]
norm = lambda e: "".join(ch for ch in (e or "").upper() if ch.isalnum())
visits = {r["voter_key"]: r for r in prod("SELECT voter_key, form_type, asdd_correct, del_status FROM ok_verification")}
db = [x for x in glob.glob(".wrangler/state/v3/d1/**/*.sqlite", recursive=True) if not x.endswith("metadata.sqlite")][0]
c = sqlite3.connect(db)
rows = []
for r in c.execute("""SELECT epic_norm, epic_id, house_raw, hh_sort, name, relation_type, relation_name, age, gender, status, status_detail, old_part, block, house_floor
                      FROM ok_electors WHERE status IN ('ASDD','DELETED') AND epic_norm <> ''"""):
    v = visits.get(r[0])
    if not v: continue
    if r[9] == "ASDD" and v["form_type"] == "asdd" and v["asdd_correct"] == "no": txt = f"ASDD → {r[10] or 'category not recorded'}"
    elif r[9] == "DELETED" and v["form_type"] == "deleted" and v["del_status"] == "wrongful": txt = f"Pre-SIR deletion (found wrongful) → {r[10] or ''}"
    else: continue
    rows.append((r[12] or "—", r[11], r[3] or "", r[2], r[13], r[4], r[5], r[6], r[7], r[8], r[1], txt))
def hn(h): return (h or "").split(",")[0].strip() or "no house number"
secs = defaultdict(list)
for x in rows: secs[(x[0], x[1])].append(x)
for k in secs: secs[k].sort(key=lambda x: (x[2], (x[4] or ""), x[5]))
n = len(rows); nd = sum(1 for x in rows if x[11].startswith("Pre-SIR")); na = n - nd
sm = ParagraphStyle("s", fontName="A", fontSize=8.6, leading=10.6); sb = ParagraphStyle("sb", parent=sm, fontName="AB")
hd = ParagraphStyle("h", parent=sm, fontName="AB", textColor=colors.white)
def P(t, st=sm): return Paragraph(str(t).replace("&", "&amp;").replace("<", "&lt;"), st)
HEAD = ["Sr.", "House number", "Floor", "Elector", "Relative", "Age / Sex", "Voter ID (EPIC)", "Status"]
W = [mm * w for w in (10, 27, 19, 52, 50, 17, 28, 74)]
stamp = datetime.datetime.now().strftime("%d %b %Y, %H:%M")
story = [Paragraph("Okhla AC 54 — electors found WRONGLY marked ASDD or wrongly deleted before the SIR", ParagraphStyle("t", fontName="AB", fontSize=14, leading=17)), Spacer(1, 2 * mm),
         Paragraph(f"<b>{n} electors</b> in {len({(x[0], x[1], x[3]) for x in rows})} households: {na} marked ASDD wrongly" + (f", {nd} pre-SIR deletion found wrongful" if nd else "") +
                   f". Listed block by block, house numbers small to large. Based on the field team's visits as of {stamp} IST.", ParagraphStyle("u", parent=sm, fontSize=9, leading=12)), Spacer(1, 4 * mm)]
first = True; sn = 0
for (block, part) in sorted(secs, key=lambda k: (k[0] == "—", k[0], k[1])):
    items = secs[(block, part)]
    data = [[P(h, hd) for h in HEAD]]; styles = []; prev = None
    for x in items:
        sn += 1
        newh = x[3] != prev
        if newh: styles.append(("LINEABOVE", (0, len(data)), (-1, len(data)), 0.7, colors.HexColor("#666666")))
        data.append([P(sn), P(hn(x[3]) if newh else "", sb), P((x[4] or "").strip()), P(x[5], sb), P(f"{x[6]}: {x[7]}" if x[7] else ""),
                     P(" / ".join(t for t in (x[8] or "", (x[9] or "")[:1]) if t)), P(x[10] or ""), P(x[11])])
        prev = x[3]
    t = Table(data, colWidths=W, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243B55")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                           ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)] + styles))
    if not first: story.append(PageBreak())
    first = False
    story += [Paragraph((f"Block {block}" if block != "—" else "No block on record") + f" — Old Part {part}", ParagraphStyle("h2", fontName="AB", fontSize=12, leading=15)),
              Paragraph(f"{len(items)} electors in {len({x[3] for x in items})} households", ParagraphStyle("c", parent=sm, textColor=colors.HexColor("#444444"))), Spacer(1, 2 * mm), t]
def footer(cv, doc):
    cv.saveState(); cv.setFont("A", 7); cv.setFillColor(colors.grey)
    cv.drawString(10 * mm, 7 * mm, f"Okhla (AC 54) · electors found wrongly marked ASDD / wrongly deleted · {stamp} IST · Confidential: contains voter personal data")
    cv.drawRightString(landscape(A4)[0] - 10 * mm, 7 * mm, f"Page {doc.page}"); cv.restoreState()
base = "okhla_import/out/Okhla_wrongly_marked_electors_by_block"
doc = BaseDocTemplate(base + ".pdf", pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm, title="Okhla electors wrongly marked ASDD / deleted")
doc.addPageTemplates([PageTemplate(frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, 0, 0, 0, 0)], onPage=footer)])
doc.build(story)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Electors"
ws.append(["Block", "Old part", "House number", "Floor", "Elector", "Relative", "Age", "Sex", "Voter ID (EPIC)", "Status"])
for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="243B55")
for (block, part) in sorted(secs, key=lambda k: (k[0] == "—", k[0], k[1])):
    for x in secs[(block, part)]:
        ws.append([block, part, hn(x[3]), (x[4] or "").strip(), x[5], f"{x[6]}: {x[7]}" if x[7] else "", x[8], (x[9] or "")[:1], x[10], x[11]])
for col, w in zip("ABCDEFGHIJ", (7, 9, 14, 10, 30, 30, 6, 5, 16, 52)): ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions; wb.save(base + ".xlsx")
print(base, n, "electors", {f"{b}/{p}": len(v) for (b, p), v in sorted(secs.items())}, "ASDD", na, "deleted", nd)
